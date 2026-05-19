"""
Extracción local de texto desde adjuntos (Camino B).

Se procesa el adjunto en el servidor (sin subirlo al LLM) y su texto se
concatena al transcript del turno con un separador estándar.

Formatos soportados:
- PDF (`application/pdf`, `.pdf`) → ``pypdf``
- DOCX (`application/vnd.openxmlformats-officedocument.wordprocessingml.document`, `.docx`)
  → ``python-docx``
- Texto plano (`text/*`, `.txt`, `.md`) → decodificación directa

Incluye:
- **Caché global** por SHA-256 del binario: el mismo PDF re-subido en turnos
  consecutivos no se re-procesa.
- **Límites duros** de tamaño por archivo y número de adjuntos por turno.
"""

from __future__ import annotations

import base64
import hashlib
import io
from dataclasses import dataclass

import structlog

logger = structlog.get_logger(__name__)

ATTACHMENT_SEPARATOR_TEMPLATE = "--- attachment: {filename} ---"

MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
"""Tamaño máximo por archivo (10 MB)."""

MAX_ATTACHMENTS_PER_TURN = 5
"""Número máximo de adjuntos en un único turno."""

SUPPORTED_IMAGE_MIME_TYPES: frozenset[str] = frozenset(
    {"image/png", "image/jpeg", "image/gif", "image/webp"},
)

_TEXT_CACHE: dict[str, str] = {}
"""Mapa SHA-256(bytes) → texto extraído. Caché simple in-process."""


def cache_size() -> int:
    return len(_TEXT_CACHE)


def reset_cache() -> None:
    _TEXT_CACHE.clear()


@dataclass(frozen=True)
class AttachmentText:
    """Resultado de extraer texto de un adjunto."""

    filename: str
    text: str
    error: str | None = None


@dataclass(frozen=True)
class AttachmentImage:
    """Imagen lista para enviar al LLM (Camino A, multimodal)."""

    filename: str
    mime_type: str
    base64_data: str
    error: str | None = None


def _guess_image_mime(filename: str, content_type: str | None) -> str | None:
    ct = (content_type or "").lower()
    if ct in SUPPORTED_IMAGE_MIME_TYPES:
        return ct
    name = (filename or "").lower()
    if name.endswith(".png"):
        return "image/png"
    if name.endswith((".jpg", ".jpeg")):
        return "image/jpeg"
    if name.endswith(".gif"):
        return "image/gif"
    if name.endswith(".webp"):
        return "image/webp"
    return None


def is_image_attachment(filename: str, content_type: str | None) -> bool:
    return _guess_image_mime(filename, content_type) is not None


def encode_image_attachment(
    *,
    filename: str,
    content_type: str | None,
    data: bytes,
) -> AttachmentImage:
    """Convierte una imagen a su representación multimodal (base64)."""
    if len(data) > MAX_ATTACHMENT_BYTES:
        return AttachmentImage(
            filename=filename,
            mime_type="",
            base64_data="",
            error=(
                f"Imagen demasiado grande: {len(data)} bytes "
                f"(límite: {MAX_ATTACHMENT_BYTES})."
            ),
        )

    mime = _guess_image_mime(filename, content_type)
    if mime is None:
        return AttachmentImage(
            filename=filename,
            mime_type="",
            base64_data="",
            error=f"Tipo de imagen no soportado: {content_type!r}",
        )

    encoded = base64.b64encode(data).decode("ascii")
    return AttachmentImage(filename=filename, mime_type=mime, base64_data=encoded)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text:
            parts.append(page_text)
    return "\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text:
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_plain(data: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


def extract_attachment_text(
    *,
    filename: str,
    content_type: str | None,
    data: bytes,
) -> AttachmentText:
    """Extrae texto plano de un adjunto, con caché por SHA-256.

    El mismo binario (mismo hash) reutiliza el texto extraído previamente,
    incluso si llega con otro filename o content_type.
    """
    if len(data) > MAX_ATTACHMENT_BYTES:
        return AttachmentText(
            filename=filename,
            text="",
            error=(
                f"Archivo demasiado grande: {len(data)} bytes "
                f"(límite: {MAX_ATTACHMENT_BYTES} bytes)."
            ),
        )

    name_lower = (filename or "").lower()
    ct = (content_type or "").lower()
    digest = hashlib.sha256(data).hexdigest()

    cached = _TEXT_CACHE.get(digest)
    if cached is not None:
        logger.debug("attachment_cache_hit", filename=filename, sha256=digest[:12])
        return AttachmentText(filename=filename, text=cached)

    try:
        if ct == "application/pdf" or name_lower.endswith(".pdf"):
            text = _extract_pdf(data)
        elif (
            ct
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or name_lower.endswith(".docx")
        ):
            text = _extract_docx(data)
        elif ct.startswith("text/") or name_lower.endswith((".txt", ".md")):
            text = _extract_plain(data)
        else:
            return AttachmentText(
                filename=filename,
                text="",
                error=f"Tipo no soportado: content_type={ct!r}",
            )
    except Exception as exc:
        logger.warning(
            "attachment_extract_failed",
            filename=filename,
            content_type=ct,
            error=str(exc),
        )
        return AttachmentText(filename=filename, text="", error=str(exc))

    _TEXT_CACHE[digest] = text
    return AttachmentText(filename=filename, text=text)


def build_augmented_transcript(
    transcript: str,
    attachments: list[AttachmentText],
) -> str:
    """Concatena el transcript del usuario con el texto extraído de los adjuntos."""
    parts: list[str] = [transcript.strip()] if transcript.strip() else []
    for att in attachments:
        if not att.text:
            continue
        parts.append(ATTACHMENT_SEPARATOR_TEMPLATE.format(filename=att.filename))
        parts.append(att.text)
    return "\n\n".join(parts).strip()
